"""Document processing: scan hook, text extraction, chunking, embedding."""

from __future__ import annotations

import hashlib
import io

from sqlalchemy import select
from sqlalchemy.orm import Session

from app.ai.providers import get_embedder
from app.core.enums import DocumentStatus, ScanStatus
from app.models.documents import (
    Document,
    DocumentChunk,
    DocumentText,
    DocumentVersion,
)
from app.services.storage import get_storage


def scan_clean(data: bytes) -> bool:
    """Malware scan hook. Local/dev: trivial allow (no EICAR signature)."""
    return b"EICAR-STANDARD-ANTIVIRUS-TEST-FILE" not in data


def extract_text(filename: str, mime: str | None, data: bytes) -> str:
    name = (filename or "").lower()
    try:
        if name.endswith(".pdf") or (mime and "pdf" in mime):
            from pypdf import PdfReader

            reader = PdfReader(io.BytesIO(data))
            return "\n".join((page.extract_text() or "") for page in reader.pages)
        if name.endswith(".docx"):
            import docx

            doc = docx.Document(io.BytesIO(data))
            return "\n".join(p.text for p in doc.paragraphs)
        if name.endswith((".xlsx", ".xlsm")):
            import openpyxl

            wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
            out = []
            for ws in wb.worksheets:
                for row in ws.iter_rows(values_only=True):
                    out.append(" ".join(str(c) for c in row if c is not None))
            return "\n".join(out)
    except Exception:  # noqa: BLE001 - extraction is best-effort
        pass
    try:
        return data.decode("utf-8", errors="ignore")
    except Exception:  # noqa: BLE001
        return ""


def chunk_text(text: str, size: int = 600, overlap: int = 80) -> list[str]:
    text = text.strip()
    if not text:
        return []
    paras = [p.strip() for p in text.split("\n") if p.strip()]
    chunks: list[str] = []
    buf = ""
    for p in paras:
        if len(buf) + len(p) + 1 <= size:
            buf = f"{buf}\n{p}" if buf else p
        else:
            if buf:
                chunks.append(buf)
            buf = (buf[-overlap:] + "\n" + p) if buf else p
            if len(buf) > size:
                # hard-split very long paragraphs
                for i in range(0, len(p), size):
                    chunks.append(p[i : i + size])
                buf = ""
    if buf:
        chunks.append(buf)
    return chunks


def _scope_hash(doc: Document) -> str:
    raw = f"{doc.organization_id}|{doc.classification}|{doc.entity_type}|{doc.entity_id}"
    return hashlib.sha256(raw.encode()).hexdigest()[:32]


def process_version(db: Session, version: DocumentVersion) -> None:
    """Run the full pipeline for a document version (scan → extract → embed)."""
    doc = db.get(Document, version.document_id)
    if not doc:
        return
    data = get_storage().read(version.storage_key)

    # 1. Scan
    if not scan_clean(data):
        version.scan_status = ScanStatus.INFECTED
        doc.status = DocumentStatus.QUARANTINED
        db.flush()
        return
    version.scan_status = ScanStatus.CLEAN
    doc.status = DocumentStatus.PROCESSING

    # 2. Extract
    text = extract_text(version.filename, version.mime_type, data)
    db.add(
        DocumentText(
            organization_id=doc.organization_id,
            document_version_id=version.id,
            content=text[:1_000_000],
        )
    )

    # 3. Chunk + embed (mark previous chunks of this doc as not current)
    for old in db.scalars(select(DocumentChunk).where(DocumentChunk.document_id == doc.id)).all():
        old.is_current = False
    chunks = chunk_text(text)
    if chunks:
        vectors = get_embedder().embed(chunks)
        scope = _scope_hash(doc)
        for i, (c, vec) in enumerate(zip(chunks, vectors, strict=False)):
            db.add(
                DocumentChunk(
                    organization_id=doc.organization_id,
                    document_id=doc.id,
                    document_version_id=version.id,
                    chunk_no=i,
                    text=c,
                    token_count=max(1, len(c) // 4),
                    classification=doc.classification,
                    access_scope_hash=scope,
                    is_current=True,
                    embedding=vec,
                )
            )

    version.extraction_status = "DONE"
    doc.status = DocumentStatus.AVAILABLE
    db.flush()
